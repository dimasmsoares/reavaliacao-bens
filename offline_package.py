"""
Avaliação offline: pacote para download e importação de volta.

O servidor baixa um .zip com uma planilha (uma linha por grupo de bens idênticos
pendentes) e uma pasta prints/<código do grupo>/ para cada grupo. Avalia fora do
sistema e envia o .zip (ou só a planilha) de volta. A importação tem duas etapas:
parse_upload() gera uma prévia (sem gravar nada) e apply_import() grava as linhas
válidas, com as mesmas regras e a mesma propagação da tela de avaliação.
"""
import io
import json
import os
import re
import shutil
import time
import unicodedata
import uuid
import zipfile
from datetime import datetime, date, timedelta
from urllib.parse import quote

import openpyxl
import requests as http_requests
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from PIL import Image

import database as db
from pdf_report import REFERENCIA_PRECO_LABELS

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
SCREENSHOTS_DIR = os.path.join(BASE_DIR, 'screenshots')
UPLOADS_DIR = os.path.join(BASE_DIR, 'offline_uploads')

PACKAGE_VERSION = 1
SHEET_NAME = 'Avaliação'
CONTROL_SHEET = '_controle'
N_PRICES = 5
IMAGE_EXTS = ('.jpg', '.jpeg', '.png', '.webp', '.bmp', '.gif')
MAX_IMAGE_BYTES = 30 * 1024 * 1024
MAX_TOTAL_UNCOMPRESSED = 2 * 1024 * 1024 * 1024
SCREENSHOT_MAX_W = 1280      # mesmo padrão da compressão feita no navegador (app.js)
SCREENSHOT_QUALITY = 85
UPLOAD_TTL = timedelta(days=1)
TUTORIAL_PATH = os.path.join(BASE_DIR, 'static', 'tutorial_offline.html')

# Colunas da planilha: (chave, cabeçalho, largura, editável)
COLUMNS = [
    ('codigo',      'Código',               9,  False),
    ('id',          'ID',                   6,  False),   # oculta: bem representante
    ('material',    'Material',             38, False),
    ('marca',       'Marca',                16, False),
    ('modelo',      'Modelo',               22, False),
    ('tipo',        'Tipo',                 11, False),
    ('qtd',         'Qtd. bens',            9,  False),
    ('ml',          'Mercado Livre',        14, False),
    ('gshop',       'Google Shopping',      16, False),
    ('metodologia', 'Metodologia',          12, True),
] + [
    (f'preco{i}', f'Preço {i}', 12, True) for i in range(1, N_PRICES + 1)
] + [
    ('valor',       'Valor de mercado',     16, True),
    ('referencia',  'Referência do valor',  18, True),
    ('observacao',  'Observação',           30, True),
    ('vc',          'Valor contábil',       14, False),
    ('tombamento',  'Data de tombamento',   13, False),
    ('planilha',    'Planilha',             30, False),
    ('nrps',        'NRPs (exemplos)',      30, False),
    ('prompt_ia',   'Pedido para IA (copie e cole)', 40, False),
]
COL = {key: idx for idx, (key, *_rest) in enumerate(COLUMNS, start=1)}
HEADER_TO_KEY = {header: key for key, header, *_rest in COLUMNS}

AI_PROMPT = ('Pesquise o valor médio de mercado atual no Brasil (em reais) para o seguinte '
             'bem usado, com base em anúncios reais de venda (não doação): "{desc}". '
             'Considere o estado de conservação típico de um bem em uso há alguns anos. '
             'Responda com uma faixa de preço e uma estimativa de valor médio em R$, citando '
             'de forma resumida em quais anúncios/fontes se baseou.')

LEIA_ME = """\
AVALIAÇÃO OFFLINE — REAVALIAÇÃO DE BENS MÓVEIS
Servidor: {user}   ·   Gerado em: {data}   ·   {n} grupo(s) pendente(s)

>> Passo a passo completo, com exemplos: abra o arquivo TUTORIAL.html
   (funciona no navegador, mesmo sem internet).

1. Extraia este .zip em uma pasta do seu computador.

2. Abra "avaliacao.xlsx". Cada linha é um GRUPO de bens idênticos: avaliar a
   linha avalia todos os bens do grupo (coluna "Qtd. bens").

3. Para cada grupo que avaliar, preencha as colunas amarelas:
   - Metodologia: M1 (pesquisa de mercado), M2 (acervo patrimonial) ou
     M3 (correção pelo IPCA).
   - M1/M2: informe de 1 a 5 preços encontrados. O "Valor de mercado" é a média,
     calculada automaticamente; você pode digitar outro valor por cima.
   - M3: basta escolher "M3". O sistema busca o IPCA e calcula o valor na
     importação (precisa de valor contábil e data de tombamento).
   - Referência do valor (opcional, M1/M2): novo/usado e igual/similar.
   - Observação (opcional).
   Os links "Mercado Livre" e "Google Shopping" abrem a pesquisa do bem. A última
   coluna traz um pedido pronto para colar no Claude/ChatGPT/Gemini.

4. Prints: salve as imagens da pesquisa (2 ou 3 por grupo) dentro da pasta do
   grupo em "prints/". Ex.: prints/G0012 - CADEIRA .../print1.png
   Não cole as imagens dentro da planilha — elas seriam ignoradas.

5. Para enviar, compacte a pasta inteira em um .zip (botão direito > Enviar para
   > Pasta compactada) e envie em "Offline > Importar" no sistema. Você pode
   enviar só a planilha, se ainda não tiver prints.

6. Pode importar em várias etapas: linhas em branco são ignoradas e linhas já
   importadas não são gravadas de novo. Antes de gravar, o sistema mostra uma
   prévia com o que será salvo, avisos e conflitos.

Não altere as colunas cinzas nem os códigos dos grupos.
"""


# ── Utilidades ──────────────────────────────────────────────────────────────

def _strip_codigo(material):
    return re.sub(r'\s*\(\d+\)\s*$', '', material or '').strip()


def _planilha_curta(planilha):
    return planilha.split(' - ', 1)[1] if ' - ' in planilha else planilha


def _ascii(text):
    return unicodedata.normalize('NFKD', text).encode('ascii', 'ignore').decode('ascii')


def _folder_name(code, desc):
    """Nome da pasta de prints do grupo: código + descrição curta, só ASCII (o
    descompactador nativo do Windows não respeita nomes UTF-8 em zip)."""
    clean = re.sub(r'[<>:"/\\|?*\x00-\x1f]+', ' ', _ascii(desc))
    clean = re.sub(r'\s+', ' ', clean).strip()[:50].rstrip(' .')
    return f'{code} - {clean}' if clean else code


def _norm_label(text):
    return _ascii(str(text)).strip().lower()


REFERENCIA_BY_TEXT = {}
for _key, _label in REFERENCIA_PRECO_LABELS.items():
    REFERENCIA_BY_TEXT[_norm_label(_label)] = _key
    REFERENCIA_BY_TEXT[_norm_label(_key)] = _key
    REFERENCIA_BY_TEXT[_norm_label(_key.replace('_', ' '))] = _key


def _parse_number(v):
    """Número de uma célula: aceita numérico ou texto em formato brasileiro."""
    if v is None or v == '':
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace('R$', '').replace(' ', '').replace('\xa0', '')
    if not s:
        return None
    if ',' in s:
        s = s.replace('.', '').replace(',', '.')
    try:
        return float(s)
    except ValueError:
        raise ValueError(f'"{v}" não é um número válido')


def _parse_date_br(s):
    m = re.search(r'(\d{1,2})/(\d{1,2})/(\d{4})', s or '')
    if not m:
        return None
    try:
        return date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
    except ValueError:
        return None


def _code_from_path(path):
    """Código do grupo (G0001) a partir do caminho de um arquivo dentro do zip."""
    for part in path.replace('\\', '/').split('/'):
        m = re.match(r'^(g\d{4,})(?![0-9])', part.strip(), re.IGNORECASE)
        if m:
            return m.group(1).upper()
    return None


def _is_junk(name):
    base = os.path.basename(name)
    return (name.endswith('/') or '__MACOSX' in name or base.startswith('~$')
            or base.startswith('._') or base.lower() in ('thumbs.db', 'desktop.ini'))


# ── IPCA ────────────────────────────────────────────────────────────────────

def ipca_acumulado_por_data(datas):
    """IPCA acumulado (%) desde cada data até hoje, com uma única consulta à série 433
    do BCB. Mesma regra de /api/ipca: entram os meses com data >= data inicial.
    Retorna {date: percentual}. Lança exceção se a API falhar."""
    datas = sorted(set(datas))
    if not datas:
        return {}
    url = ('https://api.bcb.gov.br/dados/serie/bcdata.sgs.433/dados'
           f'?formato=json&dataInicial={datas[0].strftime("%d/%m/%Y")}'
           f'&dataFinal={date.today().strftime("%d/%m/%Y")}')
    # A API do BCB devolve 502 de forma intermitente: tenta algumas vezes.
    for attempt in range(3):
        try:
            resp = http_requests.get(url, timeout=20, verify=False)
            resp.raise_for_status()
            break
        except http_requests.RequestException:
            if attempt == 2:
                raise
            time.sleep(1.5 * (attempt + 1))
    serie = [(datetime.strptime(e['data'], '%d/%m/%Y').date(),
              float(e['valor'].replace(',', '.'))) for e in resp.json()]
    result = {}
    for d in datas:
        acc = 1.0
        for dt, v in serie:
            if dt >= d:
                acc *= (1 + v / 100)
        result[d] = round((acc - 1) * 100, 4)
    return result


# ── Geração do pacote ───────────────────────────────────────────────────────

def build_package(user_id, user_name):
    """Gera o .zip do pacote offline e marca os bens como em avaliação offline.
    Retorna (bytes, nome_do_arquivo, n_grupos); n_grupos = 0 se não há pendências."""
    groups = db.get_pending_groups_for_user(user_id)
    if not groups:
        return None, None, 0
    groups.sort(key=lambda g: (g['rep']['planilha'], _strip_codigo(g['rep']['material']),
                               g['rep']['marca'] or '', g['rep']['modelo'] or ''))

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = SHEET_NAME

    header_fill = PatternFill('solid', fgColor='1F4E78')
    input_header_fill = PatternFill('solid', fgColor='BF8F00')
    input_fill = PatternFill('solid', fgColor='FFF2CC')
    locked_fill = PatternFill('solid', fgColor='F2F2F2')
    header_font = Font(bold=True, color='FFFFFF')
    unlocked = Protection(locked=False)

    for key, header, width, editable in COLUMNS:
        c = ws.cell(row=1, column=COL[key], value=header)
        c.font = header_font
        c.fill = input_header_fill if editable else header_fill
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        ws.column_dimensions[get_column_letter(COL[key])].width = width
    ws.row_dimensions[1].height = 32
    ws.column_dimensions[get_column_letter(COL['id'])].hidden = True

    folders = []
    for i, g in enumerate(groups, start=1):
        row = i + 1
        a = g['rep']
        code = f'G{i:04d}'
        desc = _strip_codigo(a['material'])
        if a['marca']:
            desc += f" - {a['marca']}"
        if a['modelo']:
            desc += f" {a['modelo']}"
        q = quote(desc)
        nrps = g['nrps']
        values = {
            'codigo': code,
            'id': a['id'],
            'material': _strip_codigo(a['material']),
            'marca': a['marca'],
            'modelo': a['modelo'],
            'tipo': a['tipo'] or 'Principal',
            'qtd': len(g['asset_ids']),
            'vc': a['valor_contabil'],
            'tombamento': a['data_tombamento'],
            'planilha': ' / '.join(_planilha_curta(p) for p in g['planilhas']),
            'nrps': ', '.join(nrps[:10]) + (f' … (+{len(nrps) - 10})' if len(nrps) > 10 else ''),
            'prompt_ia': AI_PROMPT.format(desc=desc),
        }
        for key, value in values.items():
            ws.cell(row=row, column=COL[key], value=value)
        for key, label, url in (('ml', 'Mercado Livre', f'https://lista.mercadolivre.com.br/{q}'),
                                ('gshop', 'Google Shopping',
                                 f'https://www.google.com/search?q={q}+pre%C3%A7o&tbm=shop')):
            c = ws.cell(row=row, column=COL[key], value=label)
            c.hyperlink = url
            c.style = 'Hyperlink'

        p1 = get_column_letter(COL['preco1'])
        pn = get_column_letter(COL[f'preco{N_PRICES}'])
        ws.cell(row=row, column=COL['valor'],
                value=f'=IF(COUNT({p1}{row}:{pn}{row})=0,"",ROUND(AVERAGE({p1}{row}:{pn}{row}),2))')

        for key, _h, _w, editable in COLUMNS:
            c = ws.cell(row=row, column=COL[key])
            if editable:
                c.protection = unlocked
                c.fill = input_fill
            elif key not in ('ml', 'gshop'):
                c.fill = locked_fill
            if key in ('vc', 'valor') or key.startswith('preco'):
                c.number_format = '"R$" #,##0.00'
        folders.append(_folder_name(code, desc))

    last = len(groups) + 1
    dv_met = DataValidation(type='list', formula1='"M1,M2,M3"', allow_blank=True,
                            showErrorMessage=True, errorTitle='Metodologia',
                            error='Escolha M1, M2 ou M3.')
    dv_ref = DataValidation(type='list',
                            formula1='"' + ','.join(REFERENCIA_PRECO_LABELS.values()) + '"',
                            allow_blank=True, showErrorMessage=True,
                            errorTitle='Referência do valor', error='Escolha uma das opções da lista.')
    dv_price = DataValidation(type='decimal', operator='greaterThan', formula1='0',
                              allow_blank=True, showErrorMessage=True,
                              errorTitle='Preço', error='Informe um valor numérico maior que zero.')
    for dv in (dv_met, dv_ref, dv_price):
        ws.add_data_validation(dv)
    met_col = get_column_letter(COL['metodologia'])
    ref_col = get_column_letter(COL['referencia'])
    dv_met.add(f'{met_col}2:{met_col}{last}')
    dv_ref.add(f'{ref_col}2:{ref_col}{last}')
    dv_price.add(f'{get_column_letter(COL["preco1"])}2:{get_column_letter(COL["valor"])}{last}')

    ws.freeze_panes = ws.cell(row=2, column=COL['marca'])
    ws.auto_filter.ref = f'A1:{get_column_letter(len(COLUMNS))}{last}'
    # Proteção sem senha: evita alterar por engano as colunas de identificação,
    # mantendo filtro e ajuste de largura/altura liberados.
    ws.protection.sheet = True
    ws.protection.autoFilter = False
    ws.protection.formatColumns = False
    ws.protection.formatRows = False

    now = datetime.now()
    ctl = wb.create_sheet(CONTROL_SHEET)
    for r, (k, v) in enumerate([('versao', PACKAGE_VERSION), ('user_id', user_id),
                                ('servidor', user_name),
                                ('gerado_em', now.strftime('%Y-%m-%d %H:%M:%S')),
                                ('grupos', len(groups))], start=1):
        ctl.cell(row=r, column=1, value=k)
        ctl.cell(row=r, column=2, value=v)
    ctl.sheet_state = 'veryHidden'

    xlsx = io.BytesIO()
    wb.save(xlsx)

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, 'w', zipfile.ZIP_DEFLATED) as zf:
        zf.writestr('avaliacao.xlsx', xlsx.getvalue())
        zf.writestr('LEIA-ME.txt', LEIA_ME.format(
            user=user_name, data=now.strftime('%d/%m/%Y %H:%M'), n=len(groups)
        ).replace('\n', '\r\n').encode('utf-8-sig'))
        if os.path.isfile(TUTORIAL_PATH):
            zf.write(TUTORIAL_PATH, 'TUTORIAL.html')
        for folder in folders:
            zf.writestr(zipfile.ZipInfo(f'prints/{folder}/'), b'')

    db.set_offline_checkout(user_id, [aid for g in groups for aid in g['asset_ids']])

    safe = re.sub(r'[^\w\-]+', '_', _ascii(user_name))
    filename = f"reavaliacao_offline_{safe}_{now.strftime('%Y%m%d_%H%M')}.zip"
    return buf.getvalue(), filename, len(groups)


# ── Importação: prévia ──────────────────────────────────────────────────────

class OfflineImportError(Exception):
    """Erro que impede a leitura do arquivo enviado (mensagem para o usuário)."""


def cleanup_old_uploads():
    if not os.path.isdir(UPLOADS_DIR):
        return
    limit = datetime.now() - UPLOAD_TTL
    for name in os.listdir(UPLOADS_DIR):
        path = os.path.join(UPLOADS_DIR, name)
        try:
            if datetime.fromtimestamp(os.path.getmtime(path)) < limit:
                shutil.rmtree(path, ignore_errors=True)
        except OSError:
            pass


def save_upload(file_storage):
    """Salva o arquivo enviado em offline_uploads/<token>/ e devolve o token."""
    name = (file_storage.filename or '').lower()
    if name.endswith('.zip'):
        ext = 'zip'
    elif name.endswith('.xlsx'):
        ext = 'xlsx'
    else:
        raise OfflineImportError('Envie o pacote .zip ou a planilha .xlsx.')
    cleanup_old_uploads()
    token = uuid.uuid4().hex
    folder = os.path.join(UPLOADS_DIR, token)
    os.makedirs(folder)
    file_storage.save(os.path.join(folder, f'upload.{ext}'))
    return token


def _upload_path(token):
    if not re.fullmatch(r'[0-9a-f]{32}', token or ''):
        return None, None
    folder = os.path.join(UPLOADS_DIR, token)
    for ext in ('zip', 'xlsx'):
        p = os.path.join(folder, f'upload.{ext}')
        if os.path.isfile(p):
            return folder, p
    return None, None


def discard_upload(token):
    folder, _ = _upload_path(token)
    if folder:
        shutil.rmtree(folder, ignore_errors=True)


def _open_upload(path):
    """Retorna (bytes da planilha, {código: [membros de imagem]}, zipfile|None, avisos)."""
    if path.endswith('.xlsx'):
        with open(path, 'rb') as f:
            return f.read(), {}, None, []
    try:
        zf = zipfile.ZipFile(path)
    except zipfile.BadZipFile:
        raise OfflineImportError('O arquivo enviado não é um .zip válido.')
    infos = [i for i in zf.infolist() if not _is_junk(i.filename)]
    if sum(i.file_size for i in infos) > MAX_TOTAL_UNCOMPRESSED:
        raise OfflineImportError('O conteúdo do .zip é grande demais (mais de 2 GB descompactado).')
    xlsx = [i for i in infos if i.filename.lower().endswith('.xlsx')]
    if not xlsx:
        raise OfflineImportError('Não encontrei a planilha (.xlsx) dentro do .zip.')
    if len(xlsx) > 1:
        raise OfflineImportError('O .zip tem mais de uma planilha .xlsx: '
                           + ', '.join(i.filename for i in xlsx) + '. Deixe apenas uma.')
    warnings = []
    prints = {}
    for i in infos:
        if not i.filename.lower().endswith(IMAGE_EXTS):
            continue
        code = _code_from_path(i.filename)
        if not code:
            warnings.append(f'Imagem fora da pasta de um grupo, ignorada: {i.filename}')
        elif i.file_size > MAX_IMAGE_BYTES:
            warnings.append(f'Imagem maior que 30 MB, ignorada: {i.filename}')
        else:
            prints.setdefault(code, []).append(i.filename)
    for code in prints:
        prints[code].sort()
    return zf.read(xlsx[0]), prints, zf, warnings


def _valid_image(zf, member):
    try:
        with Image.open(io.BytesIO(zf.read(member))) as img:
            img.verify()
        return True
    except Exception:
        return False


def parse_upload(token, user_id):
    """Lê o arquivo enviado e monta a prévia (nada é gravado). A prévia fica salva em
    offline_uploads/<token>/preview.json para a confirmação."""
    folder, path = _upload_path(token)
    if not path:
        raise OfflineImportError('Arquivo de importação não encontrado. Envie novamente.')
    xlsx_bytes, prints, zf, file_warnings = _open_upload(path)
    try:
        try:
            wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), data_only=True)
        except Exception:
            raise OfflineImportError('Não foi possível abrir a planilha. Verifique se é o arquivo '
                               '"avaliacao.xlsx" do pacote, salvo no formato .xlsx.')
        if CONTROL_SHEET not in wb.sheetnames or SHEET_NAME not in wb.sheetnames:
            raise OfflineImportError('Esta planilha não é um pacote de avaliação offline do sistema.')
        ctl = {wb[CONTROL_SHEET].cell(row=r, column=1).value:
               wb[CONTROL_SHEET].cell(row=r, column=2).value for r in range(1, 10)}
        if ctl.get('user_id') != user_id:
            raise OfflineImportError(f'Este pacote foi gerado para outro servidor '
                               f'({ctl.get("servidor") or "desconhecido"}).')
        ws = wb[SHEET_NAME]
        if getattr(ws, '_images', None):
            file_warnings.append(f'{len(ws._images)} imagem(ns) colada(s) dentro da planilha '
                                 f'foram ignoradas. Salve os prints nas pastas do .zip.')

        cols = {}
        for c in ws[1]:
            key = HEADER_TO_KEY.get(str(c.value).strip()) if c.value else None
            if key:
                cols[key] = c.column
        missing = [h for k, h, *_r in COLUMNS if k not in cols]
        if missing:
            raise OfflineImportError('Colunas ausentes na planilha: ' + ', '.join(missing))

        rows, skipped_blank = [], 0
        codes_seen = set()
        for r in range(2, ws.max_row + 1):
            get = lambda k: ws.cell(row=r, column=cols[k]).value
            code = str(get('codigo') or '').strip().upper()
            if not code:
                continue
            item = _parse_row(r, code, get, user_id, prints.get(code, []))
            codes_seen.add(code)
            if item is None:
                skipped_blank += 1
            else:
                rows.append(item)

        for code in sorted(set(prints) - codes_seen):
            file_warnings.append(f'Pasta de prints "{code}" não corresponde a nenhum grupo da planilha.')

        # Prints ilegíveis
        if zf:
            for item in rows:
                good = [m for m in item['prints'] if _valid_image(zf, m)]
                for m in set(item['prints']) - set(good):
                    item['warnings'].append(f'Arquivo ignorado (não é uma imagem válida): {os.path.basename(m)}')
                item['prints'] = good

        _resolve_ipca(rows)
        for item in rows:
            _finalize_status(item)
    finally:
        if zf:
            zf.close()

    preview = {
        'user_id': user_id,
        'gerado_em': str(ctl.get('gerado_em') or ''),
        'file_warnings': file_warnings,
        'skipped_blank': skipped_blank,
        'rows': rows,
    }
    with open(os.path.join(folder, 'preview.json'), 'w', encoding='utf-8') as f:
        json.dump(preview, f, ensure_ascii=False)
    return preview


def _parse_row(r, code, get, user_id, prints):
    """Interpreta uma linha. Retorna None se a linha está em branco (não avaliada)."""
    errors, warnings = [], []
    met = str(get('metodologia') or '').strip().upper()
    raw_prices = [get(f'preco{i}') for i in range(1, N_PRICES + 1)]
    referencia_raw = get('referencia')
    observacao = str(get('observacao') or '').strip() or None
    valor_raw = get('valor')

    prices = []
    for i, v in enumerate(raw_prices, start=1):
        try:
            n = _parse_number(v)
        except ValueError as e:
            errors.append(f'Preço {i}: {e}.')
            continue
        if n is not None:
            if n <= 0:
                errors.append(f'Preço {i} deve ser maior que zero.')
            else:
                prices.append(n)
    try:
        valor = _parse_number(valor_raw)
    except ValueError as e:
        valor = None
        errors.append(f'Valor de mercado: {e}.')

    filled = met or prices or referencia_raw or observacao or prints or (valor and valor > 0)
    if not filled and not errors:
        return None

    try:
        asset_id = int(get('id'))
    except (TypeError, ValueError):
        asset_id = None
    asset = db.get_asset(asset_id) if asset_id else None

    item = {
        'row': r, 'code': code, 'asset_id': asset_id,
        'descricao': '', 'qtd': get('qtd'),
        'metodologia': None, 'prices': prices, 'valor': None, 'media': None,
        'ipca': None, 'referencia': None, 'observacao': observacao,
        'prints': list(prints), 'errors': errors, 'warnings': warnings,
        'status': None, 'status_msg': '',
    }
    if not asset:
        errors.append('Bem não encontrado no sistema (coluna ID alterada?).')
        return item
    desc = _strip_codigo(asset['material'])
    if asset['marca']:
        desc += f" — {asset['marca']}"
    if asset['modelo']:
        desc += f" {asset['modelo']}"
    item['descricao'] = desc

    if not met:
        if prices:
            met = 'M1'
            warnings.append('Metodologia não informada; considerada M1.')
        else:
            errors.append('Informe a metodologia (M1, M2 ou M3).')
    elif met not in ('M1', 'M2', 'M3'):
        errors.append(f'Metodologia inválida: "{met}". Use M1, M2 ou M3.')
    item['metodologia'] = met or None

    if met in ('M1', 'M2'):
        if not prices:
            errors.append('Informe ao menos um preço encontrado.')
        else:
            media = round(sum(prices) / len(prices), 2)
            item['media'] = media
            if valor is not None and valor <= 0:
                errors.append('Valor de mercado deve ser maior que zero.')
            elif valor is not None and abs(valor - media) > 0.005:
                item['valor'] = valor
                warnings.append(f'Valor de mercado editado (a média dos preços é '
                                f'{_brl(media)}).')
            else:
                item['valor'] = media
        if not prints:
            warnings.append('Nenhum print anexado.')
        if referencia_raw not in (None, ''):
            ref = REFERENCIA_BY_TEXT.get(_norm_label(referencia_raw))
            if ref:
                item['referencia'] = ref
            else:
                warnings.append(f'Referência do valor "{referencia_raw}" não reconhecida; ignorada.')
    elif met == 'M3':
        if not asset['valor_contabil']:
            errors.append('Bem sem valor contábil: não é possível aplicar M3.')
        tomb = _parse_date_br(asset['data_tombamento'])
        if not tomb:
            errors.append('Bem sem data de tombamento: não é possível aplicar M3.')
        item['_tombamento'] = tomb.isoformat() if tomb else None
        item['_vc'] = asset['valor_contabil']
        if prices:
            warnings.append('Preços informados foram ignorados (M3 usa a correção pelo IPCA).')
        item['prices'] = []
        if prints:
            warnings.append('Prints ignorados (não se aplicam à M3).')

    # Situação do grupo no sistema
    state = db.get_user_group_state(asset_id, user_id)
    if not state or state['assigned'] == 0:
        item['status'] = 'conflito'
        item['status_msg'] = 'Grupo não está mais atribuído a você.'
    elif not state['pending_ids']:
        reviews = state['reviews']
        if all(rv['review_user_id'] == user_id and rv['origem'] == 'offline' for rv in reviews):
            item['status'] = 'ja_importado'
            item['status_msg'] = 'Já importado anteriormente; não será gravado de novo.'
        else:
            item['status'] = 'conflito'
            item['status_msg'] = 'Grupo já foi avaliado no sistema; esta linha será ignorada.'
    return item


def _resolve_ipca(rows):
    pending = [it for it in rows
               if it['metodologia'] == 'M3' and not it['errors'] and it['status'] is None]
    if not pending:
        return
    datas = {date.fromisoformat(it['_tombamento']) for it in pending}
    try:
        ipca = ipca_acumulado_por_data(datas)
    except Exception:
        for it in pending:
            it['errors'].append('Não foi possível consultar o IPCA no Banco Central. '
                                'Tente importar novamente mais tarde.')
        return
    for it in pending:
        pct = ipca[date.fromisoformat(it['_tombamento'])]
        it['ipca'] = pct
        it['valor'] = round(it['_vc'] * (1 + pct / 100), 2)


def _finalize_status(item):
    for k in ('_tombamento', '_vc'):
        item.pop(k, None)
    if item['status']:
        return
    if item['errors']:
        item['status'] = 'erro'
    elif item['warnings']:
        item['status'] = 'aviso'
    else:
        item['status'] = 'ok'


def _brl(v):
    return f'R$ {v:,.2f}'.replace(',', 'X').replace('.', ',').replace('X', '.')


def load_preview(token, user_id):
    folder, path = _upload_path(token)
    if not folder:
        return None
    try:
        with open(os.path.join(folder, 'preview.json'), encoding='utf-8') as f:
            preview = json.load(f)
    except (OSError, ValueError):
        return None
    return preview if preview.get('user_id') == user_id else None


# ── Importação: gravação ────────────────────────────────────────────────────

def _save_print(zf, member, anchor_id, idx, stamp):
    os.makedirs(SCREENSHOTS_DIR, exist_ok=True)
    with Image.open(io.BytesIO(zf.read(member))) as img:
        img = img.convert('RGB')
        if img.width > SCREENSHOT_MAX_W:
            h = round(img.height * SCREENSHOT_MAX_W / img.width)
            img = img.resize((SCREENSHOT_MAX_W, h), Image.LANCZOS)
        filename = f'{anchor_id}_{stamp}_off{idx}.jpg'
        img.save(os.path.join(SCREENSHOTS_DIR, filename), 'JPEG',
                 quality=SCREENSHOT_QUALITY, optimize=True)
    return filename


def apply_import(token, user_id):
    """Grava as linhas 'ok' e 'aviso' da prévia. Reconfere a situação de cada grupo
    (pode ter mudado desde a prévia). Retorna {'gravados', 'bens', 'ignorados', 'prints'}."""
    preview = load_preview(token, user_id)
    if not preview:
        raise OfflineImportError('Prévia não encontrada ou expirada. Envie o arquivo novamente.')
    folder, path = _upload_path(token)
    zf = zipfile.ZipFile(path) if path.endswith('.zip') else None
    stamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    result = {'gravados': 0, 'bens': 0, 'ignorados': 0, 'prints': 0}
    try:
        for item in preview['rows']:
            if item['status'] not in ('ok', 'aviso'):
                continue
            state = db.get_user_group_state(item['asset_id'], user_id)
            if not state or not state['pending_ids']:
                result['ignorados'] += 1
                continue
            anchor = state['pending_ids'][0]
            paths = []
            if zf and item['metodologia'] in ('M1', 'M2'):
                for idx, member in enumerate(item['prints']):
                    try:
                        paths.append(_save_print(zf, member, anchor, idx, stamp))
                    except Exception:
                        pass
            db.save_review(anchor, user_id, item['valor'],
                           screenshot_path=paths[0] if paths else None,
                           prices=item['prices'] or None,
                           screenshot_paths=paths or None,
                           observacao=item['observacao'],
                           metodologia=item['metodologia'],
                           ipca_percentual=item['ipca'],
                           referencia_preco=item['referencia'],
                           origem='offline')
            result['gravados'] += 1
            result['bens'] += len(state['pending_ids'])
            result['prints'] += len(paths)
    finally:
        if zf:
            zf.close()
    shutil.rmtree(folder, ignore_errors=True)
    return result
